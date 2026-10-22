import re
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def create_api_docx(output_path="Aegis_API_Documentation.docx"):
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles & Colors
    NAVY = RGBColor(15, 23, 42)      # #0F172A
    BLUE = RGBColor(37, 99, 235)     # #2563EB
    DARK_BLUE = RGBColor(30, 58, 138)# #1E3A8A
    GRAY = RGBColor(100, 116, 139)   # #64748B
    DARK_GRAY = RGBColor(51, 65, 85) # #334155
    CODE_COLOR = RGBColor(30, 41, 59)

    # Document Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("AEGIS DISASTER INTELLIGENCE PLATFORM")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = DARK_BLUE

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(8)
    run_sub = p_sub.add_run("Official REST & Model Context Protocol (MCP) API Reference Manual")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(13)
    run_sub.font.bold = True
    run_sub.font.color.rgb = BLUE

    # Metadata Banner
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_after = Pt(24)
    run_meta = p_meta.add_run("Version 2.0.0  |  OpenAPI 3.1.0  |  Production Release")
    run_meta.font.name = "Calibri"
    run_meta.font.size = Pt(10)
    run_meta.font.italic = True
    run_meta.font.color.rgb = GRAY

    # Helper functions
    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = DARK_BLUE
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = BLUE
        return p

    def add_h3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = DARK_GRAY
        return p

    def add_p(text, bold_prefix=None, space_after=4):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = "Calibri"
            r_pre.font.size = Pt(10)
            r_pre.font.bold = True
            r_pre.font.color.rgb = DARK_GRAY
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(10)
        run.font.color.rgb = DARK_GRAY
        return p

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = "Calibri"
            r_pre.font.size = Pt(10)
            r_pre.font.bold = True
            r_pre.font.color.rgb = DARK_BLUE
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(10)
        run.font.color.rgb = DARK_GRAY
        return p

    def add_code_block(code_text):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        set_cell_background(cell, "F8FAFC")
        set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
        
        # Border
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                <w:left w:val="single" w:sz="16" w:space="0" w:color="2563EB"/>
                <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                <w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
            </w:tcBorders>
        ''')
        tcPr.append(tcBorders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(code_text.strip())
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        run.font.color.rgb = CODE_COLOR
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def style_table(table, col_widths, headers, data):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr_cells = table.rows[0].cells
        for i, header_text in enumerate(headers):
            hdr_cells[i].text = header_text
            set_cell_background(hdr_cells[i], "1E3A8A")
            set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
            p = hdr_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(9.5)
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)

        for row_idx, row_data in enumerate(data):
            row_cells = table.add_row().cells
            bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
            for col_idx, cell_value in enumerate(row_data):
                row_cells[col_idx].text = str(cell_value)
                set_cell_background(row_cells[col_idx], bg_color)
                set_cell_margins(row_cells[col_idx], top=80, bottom=80, left=140, right=140)
                p = row_cells[col_idx].paragraphs[0]
                for r in p.runs:
                    r.font.name = "Calibri"
                    r.font.size = Pt(9)
                    r.font.color.rgb = DARK_GRAY

        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------
    # SECTION 1: ARCHITECTURE & OVERVIEW
    # -------------------------------------------------------------
    add_h1("1. Overview & System Architecture")
    add_p("The Aegis Disaster Intelligence Platform exposes a high-throughput, secure REST and Model Context Protocol (MCP) interface for flood detection, continuous depth nowcasting, emergency alert dissemination, and autonomous multi-agent response orchestration.")
    
    add_bullet(" Dual-stream SegFormer MiT-B2 neural network fusing Sentinel-1 SAR (VV/VH backscatter) and Sentinel-2 L2A optical (R, G, B, NIR) bands with cloud-cover compensation.", "Multimodal Sensor Fusion:")
    add_bullet(" 3-Scale ConvGRU + YOLOX dual-head network predicting continuous water depth across 4 time horizons (0m, +15m, +30m, +60m) at 140x faster execution than MIKE+ hydrodynamic solvers.", "4D Depth Nowcasting:")
    add_bullet(" Real-time DEM-integrated graph traversal avoiding hazardous floodwaters (>30cm) to route evacuees to verified elevated shelters.", "Time-Varying Dijkstra Routing:")
    add_bullet(" 6-agent LangGraph state machine with Human-in-the-Loop approval gates for high-impact actions.", "Multi-Agent Orchestration & HITL:")
    add_bullet(" 15 secure disaster response tools executed within cryptographic provenance sandboxes.", "Formal MCP Server:")

    add_h2("1.1 Base URLs & Endpoints")
    add_bullet(" http://localhost:8000 (or http://127.0.0.1:8000)", "Local Backend:")
    add_bullet(" https://<active-subdomain>.trycloudflare.com", "Remote Public Tunnel:")
    add_bullet(" http://localhost:8000/docs", "Interactive OpenAPI Swagger UI:")
    add_bullet(" http://localhost:8000/redoc", "Alternative ReDoc UI:")

    # -------------------------------------------------------------
    # SECTION 2: AUTHENTICATION & RBAC
    # -------------------------------------------------------------
    add_h1("2. Authentication & Role-Based Access Control")
    add_p("All administrative and operational endpoints require a JSON Web Token (JWT) in the HTTP Authorization header:")
    add_code_block("Authorization: Bearer <your_jwt_access_token>")

    table = doc.add_table(rows=1, cols=3)
    headers = ["Role Name", "Access Level", "Permitted Operations"]
    roles_data = [
        ["citizen", "Public", "Submit flood reports, view detection maps, query AI chatbot, request personal evacuation routes."],
        ["official", "Restricted", "Access admin dashboard, approve HITL actions, trigger Twilio calls, resolve complaints, broadcast alerts."],
        ["ndrf", "Emergency", "Full official access with priority field rescue telemetry and dispatch logs."],
        ["admin", "Full System", "System provisioning, diagnostic suites, key management, and cryptographic audit ledger access."]
    ]
    style_table(table, [1.2, 1.2, 4.4], headers, roles_data)

    # -------------------------------------------------------------
    # SECTION 3: ENDPOINT REFERENCE
    # -------------------------------------------------------------
    add_h1("3. Core API Modules & Endpoints")

    # Module 1
    add_h2("Module 1: Authentication & User Accounts")
    
    add_h3("POST /api/auth/register")
    add_p("Registers a new citizen or official account with salted password hashing.")
    add_p("Request Schema (JSON):", bold_prefix="Payload: ")
    add_code_block("""{
  "username": "rajat_official",
  "password": "SecurePassword123!",
  "role": "official",
  "phone": "+917678656930",
  "expo_push_token": "ExponentPushToken[xxxxxxxxxxxxxx]"
}""")
    add_p("Response (200 OK):", bold_prefix="Output: ")
    add_code_block("""{
  "status": "SUCCESS",
  "message": "User registered successfully",
  "token": "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9...",
  "role": "official",
  "username": "rajat_official"
}""")

    add_h3("POST /api/auth/login")
    add_p("Authenticates user credentials and returns a signed JWT access token.")
    add_code_block("""{
  "username": "rajat_official",
  "password": "SecurePassword123!"
}""")

    # Module 2
    add_h2("Module 2: Multimodal 4D Flood Detection & Satellite Inference")
    
    add_h3("POST /api/run-detection")
    add_p("Initiates an asynchronous 16-step detection pipeline acquiring real-time Sentinel-2 L2A optical and Sentinel-1 SAR imagery.")
    add_p("Query Parameters: lat (float, required), lon (float, required), cloud_cover (float, optional).")
    add_code_block("""// Example Request:
POST /api/run-detection?lat=25.6124&lon=85.1376&cloud_cover=12.0

// Response (200 OK):
{
  "job_id": "c2bdb31a5c4d4eb8be981b1087c30e5a"
}""")

    add_h3("GET /api/status/{job_id}")
    add_p("Polls detection progress, progressive step messages, and final inference artifacts.")
    add_code_block("""{
  "job_id": "c2bdb31a5c4d4eb8be981b1087c30e5a",
  "status": "complete",
  "current_step": "dynamic_routing",
  "result": {
    "confidence_score": 91.4,
    "classification": "Flood Inundation",
    "area_sq_km": 3.84,
    "severity": "CRITICAL",
    "severity_score": 96.8,
    "impact": {
      "population": 2450,
      "buildings": 312,
      "facilities": 4
    },
    "mask_geojson": { "type": "FeatureCollection", "features": [...] },
    "spread_24h_48h_geojson": { "type": "FeatureCollection", "features": [...] },
    "segmentation_composite_b64": "data:image/png;base64,...",
    "probability_heatmap_b64": "data:image/png;base64,...",
    "urnn_nowcast": { ... },
    "dynamic_evacuation_route": { ... }
  }
}""")

    add_h3("POST /api/upload-geotiff")
    add_p("Accepts binary multi-spectral GeoTIFF files (.tif / .tiff) for direct on-device or edge model inference.")

    add_h3("POST /api/predict-image")
    add_p("Processes ground evidence or drone RGB imagery to perform localized waterlogging segmentation.")

    # Module 3
    add_h2("Module 3: U-RNN 4D Depth Nowcasting & Hydro-Tensor Simulation")
    
    add_h3("GET /api/nowcast")
    add_p("Returns 4 continuous water depth horizons (0m, +15m, +30m, +60m), 5-factor hydrological tensor metadata, and dynamic safe routes.")
    add_p("Parameters: lat, lon, rain_mm_hr (optional override), drainage_eff (default 65.0%).")
    add_code_block("""{
  "status": "SUCCESS",
  "hydro_metadata": {
    "inst_rain_mm_hr": 28.4,
    "cumul_rain_24h_mm": 74.2,
    "mean_imperviousness": 0.58,
    "drainage_inlet_count": 86
  },
  "urnn_nowcast": {
    "architecture": "U-RNN (ConvGRU + YOLOX Dual-Head)",
    "horizons_min": [0, 15, 30, 60],
    "latency_ms": 42.5,
    "speedup_vs_mike_plus": "140x faster",
    "max_depth_0m": 0.42,
    "max_depth_60m": 0.89
  },
  "dynamic_evacuation_route": {
    "distance_km": 1.84,
    "eta_minutes": 22.5,
    "avoided_flooded_nodes_count": 14,
    "elevation_safety_score": 94.2
  }
}""")

    add_h3("POST /api/nowcast/simulate")
    add_p("Allows interactive what-if disaster simulation varying precipitation volume and storm drain blockage levels.")

    # Module 4
    add_h2("Module 4: Topographic Wetness Index (TWI) & Dynamic Evacuation Routing")
    
    add_h3("POST /api/aegis/evacuation/plan")
    add_p("Computes Time-Varying Dijkstra routing over SRTM 30m Digital Elevation Models to elevated shelters avoiding water hazards.")
    add_code_block("""{
  "origin_lat": 25.6124,
  "origin_lon": 85.1376
}""")

    add_h3("POST /api/aegis/twi/predict")
    add_p("Calculates terrain slope, flow accumulation, and natural topographic pooling susceptibility nodes.")

    add_h3("POST /api/aegis/triage/run")
    add_p("Executes automated LangGraph multi-criteria resource allocation algorithm across regional shelters and citizen clusters.")

    # Module 5
    add_h2("Module 5: Citizen Crowdsourcing & Disaster Reports")
    
    add_h3("POST /api/report-flood")
    add_p("Logs ground-truth citizen complaints with photo evidence, verifies local rainfall telemetry, and queues for emergency response.")
    add_code_block("""{
  "username": "shahzeb_citizen",
  "lat": 28.6253,
  "lon": 77.1065,
  "severity": "SEVERE",
  "description": "Street drainage choked after continuous cloudburst.",
  "location_name": "Hari Nagar Block B, West Delhi",
  "image_b64": "data:image/jpeg;base64,..."
}""")

    add_h3("GET /api/complaints/list")
    add_p("Retrieves filterable list of all citizen reports (status: pending, resolved; severity: LOW, MODERATE, HIGH, CRITICAL).")

    add_h3("POST /api/complaints/resolve")
    add_p("Allows verified officials to mark a complaint resolved upon emergency crew dispatch.")

    # Module 6
    add_h2("Module 6: Multi-Agent AI & Disaster RAG Chatbot")
    
    add_h3("POST /api/agent/chat")
    add_p("FloodRescuer emergency assistant powered by Google Gemini 2.5 Flash and vector-indexed NDMA guidelines.")
    add_code_block("""// Request:
{
  "message": "Where is the nearest safe shelter and how do I avoid deep water?",
  "lastResult": { "location": "Patna Sector 3", "area_sq_km": 3.84, "severity": "CRITICAL" }
}

// Response:
{
  "response": "Evacuate immediately toward the District Elevated Community Center located 2.45 km north-east. The Time-Varying Dijkstra routing engine has identified clear pathways avoiding the flooded Main Road underpass (depth: 0.74m). Keep emergency helpline 108 on speed dial."
}""")

    add_h3("POST /api/chat/message")
    add_p("Conversational MirEye tool-routing agent for geocoding and real-time hazard evaluation.")

    add_h3("POST /api/agent-cycle")
    add_p("Triggers autonomous execution across all 6 specialized agents in the LangGraph swarm.")

    # Module 7
    add_h2("Module 7: Emergency Telephony, Multilingual Voice & Broadcasts")
    
    add_h3("POST /api/alerts/voice")
    add_p("Dispatches an automated phone call using Twilio Voice API speaking Hindi or English emergency guidance.")
    add_code_block("""{
  "phone_number": "+917678656930",
  "message": "चेतावनी! आपके क्षेत्र में जलस्तर 0.68 मीटर तक बढ़ गया है। तुरंत सुरक्षित आश्रय की ओर प्रस्थान करें।",
  "location": "Patna District",
  "severity": "CRITICAL",
  "language": "hi",
  "max_depth_m": 0.68
}""")

    add_h3("POST /api/alerts/voice-polly")
    add_p("Synthesizes emergency audio broadcasts using Amazon Polly Neural TTS in Indian languages.")

    add_h3("POST /api/alerts/broadcast-sns")
    add_p("Mass-broadcasts SMS emergency evacuation notices via Amazon SNS.")

    add_h3("GET /api/alerts/cap")
    add_p("Exports standardized OASIS Common Alerting Protocol (CAP v1.2) XML feeds for national emergency broadcasting.")

    # Module 8
    add_h2("Module 8: Formal Model Context Protocol (MCP) Server")
    
    add_h3("GET /api/mcp/tools")
    add_p("Returns the formal catalog of 15 registered disaster response MCP tools.")

    add_h3("POST /api/mcp/execute")
    add_p("Executes an approved MCP tool in a cryptographically signed sandbox environment with idempotency protection.")
    add_code_block("""{
  "tool_name": "compute_urnn_depth_nowcast",
  "arguments": { "lat": 25.6124, "lon": 85.1376, "rainfall_mm_hr": 35.0 },
  "idempotency_key": "mcp_nowcast_job_991823"
}""")

    add_h3("GET /api/mcp/provenance")
    add_p("Returns the immutable cryptographic execution ledger of all tool invocations and caller signatures.")

    # Module 9 & 10
    add_h2("Module 9: Automated SITREP & PDF Reporting")
    add_h3("GET /api/reports/download-pdf")
    add_p("Dynamically compiles and downloads an official NDRF / SDMA formatted PDF Situation Report.")

    add_h2("Module 10: Human-in-the-Loop (HITL) Safety Gate")
    add_h3("GET /api/hitl/queue")
    add_p("Lists pending autonomous actions requiring commander approval.")
    add_h3("POST /api/hitl/approve")
    add_p("Authorizes execution of a draft action with commander identity verification.")

    # -------------------------------------------------------------
    # SECTION 4: CODE SAMPLES
    # -------------------------------------------------------------
    add_h1("4. Client Integration Code Samples")
    
    add_h2("4.1 Python Integration (requests)")
    add_code_block("""import requests

BASE_URL = "http://localhost:8000/api"

# 1. Authenticate
token = requests.post(f"{BASE_URL}/auth/login", json={
    "username": "saif", "password": "123"
}).json()["token"]

headers = {"Authorization": f"Bearer {token}"}

# 2. Trigger 16-Step Pipeline
job_id = requests.post(f"{BASE_URL}/run-detection", params={
    "lat": 25.6124, "lon": 85.1376, "cloud_cover": 10.0
}).json()["job_id"]

# 3. Poll for Completion
status = requests.get(f"{BASE_URL}/status/{job_id}").json()
print("Pipeline Status:", status["status"])""")

    add_h2("4.2 TypeScript / JavaScript Integration (fetch)")
    add_code_block("""const BASE_URL = 'http://localhost:8000/api';

async function fetchNowcast(lat: number, lon: number) {
  const res = await fetch(`${BASE_URL}/nowcast?lat=${lat}&lon=${lon}&drainage_eff=70`);
  const data = await res.json();
  if (data.status === 'SUCCESS') {
    console.log('Peak depth in 60 min:', data.urnn_nowcast.max_depth_60m);
    console.log('Safe Route ETA:', data.dynamic_evacuation_route.eta_minutes, 'min');
  }
}""")

    add_h2("4.3 cURL Integration")
    add_code_block("""curl -X POST "http://localhost:8000/api/agent/chat" \\
     -H "Content-Type: application/json" \\
     -d '{"message": "Where is the nearest shelter?"}'""")

    doc.save(output_path)
    print(f"Document successfully created at {output_path}")

if __name__ == "__main__":
    create_api_docx()
